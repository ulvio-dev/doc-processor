"""Build a small representative corpus: Dutch-language care procedures.

Deliberately synthetic (the real 600 documents are the customer's), but shaped
like them: headings, numbered procedure steps, a responsibilities table, and a
scanned variant of the same text.
"""
import io, os, sys
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak

OUT = sys.argv[1]
os.makedirs(OUT, exist_ok=True)

TITLE = "Procedure PRO-014: Opname en intake op de afdeling psychiatrie"
INTRO = (
    "Deze procedure beschrijft het verloop van een geplande opname op de afdeling "
    "volwassenenpsychiatrie, vanaf de aanmelding door de verwijzer tot en met de "
    "registratie van het behandelplan in het elektronisch patiëntendossier (EPD). "
    "De procedure is van toepassing op alle medewerkers van de zorgeenheid en "
    "vervangt de vorige versie van 12 maart 2024."
)
STEPS = [
    ("Aanmelding", "De verwijzer bezorgt het aanmeldingsformulier en de meest recente medische verslagen aan het onthaal. Het onthaal registreert de aanmelding binnen 24 uur in het EPD en kent een dossiernummer toe."),
    ("Voorbereidend overleg", "De behandelend psychiater beoordeelt de aanmelding op indicatie en dringendheid. Bij twijfel over de indicatie wordt de casus besproken op het wekelijkse intakeoverleg van dinsdag."),
    ("Intakegesprek", "Het intakegesprek gebeurt door een psychiater en een verpleegkundige. Er wordt gepeild naar de hulpvraag, de psychiatrische voorgeschiedenis, de medicatie, het middelengebruik en het sociale netwerk. Het risico op suicidaliteit wordt systematisch bevraagd volgens de richtlijn RIS-002."),
    ("Somatische screening", "Binnen 48 uur na opname voert de arts een somatisch onderzoek uit, aangevuld met een bloedafname en een ECG bij patiënten ouder dan 50 jaar of bij gebruik van antipsychotica."),
    ("Medicatieanamnese", "De ziekenhuisapotheker of de referentieverpleegkundige medicatie stelt het thuismedicatieschema op in overleg met de huisarts en de officina-apotheker. Afwijkingen worden gemeld aan de behandelend arts."),
    ("Opstellen behandelplan", "Binnen vijf werkdagen wordt een behandelplan opgesteld met de patiënt. Het plan bevat doelstellingen, de voorziene therapieën, de betrokken disciplines en een geplande evaluatiedatum."),
    ("Registratie en opvolging", "Het behandelplan wordt geregistreerd in het EPD onder de rubriek Behandeling. De evaluatie gebeurt minimaal tweewekelijks op het multidisciplinair teamoverleg."),
]
TABLE = [
    ["Stap", "Verantwoordelijke", "Termijn", "Registratie"],
    ["Aanmelding", "Onthaalmedewerker", "24 uur", "EPD – module Aanmelding"],
    ["Intakegesprek", "Psychiater + verpleegkundige", "3 werkdagen", "EPD – Intakeverslag"],
    ["Somatische screening", "Arts", "48 uur", "EPD – Somatiek"],
    ["Medicatieanamnese", "Apotheker", "48 uur", "Medicatieschema"],
    ["Behandelplan", "Multidisciplinair team", "5 werkdagen", "EPD – Behandeling"],
]
REFS = [
    "KB van 25 april 2002 betreffende de vaststelling en de vereffening van het budget van financiële middelen van de ziekenhuizen.",
    "Vlaams decreet betreffende de geestelijke gezondheidszorg, gecoordineerde versie.",
    "Interne richtlijn RIS-002: inschatting van suicidaliteit bij opname.",
    "Interne richtlijn MED-007: medicatiereconciliatie bij opname en ontslag.",
]

def build_pdf(path, pages=6):
    doc = SimpleDocTemplate(path, pagesize=A4, topMargin=20*mm, bottomMargin=20*mm)
    s = getSampleStyleSheet()
    story = [Paragraph(TITLE, s["Title"]), Paragraph(INTRO, s["BodyText"]), Spacer(1, 6*mm)]
    for rep in range(pages):
        story.append(Paragraph(f"{rep+1}. Werkwijze (deel {rep+1})", s["Heading1"]))
        for i, (h, body) in enumerate(STEPS, 1):
            story.append(Paragraph(f"{rep+1}.{i} {h}", s["Heading2"]))
            story.append(Paragraph(body, s["BodyText"]))
        story.append(Spacer(1, 4*mm))
        story.append(Paragraph("Verantwoordelijkheden", s["Heading2"]))
        t = Table(TABLE, colWidths=[38*mm, 45*mm, 25*mm, 50*mm])
        t.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(t)
        story.append(Paragraph("Referenties", s["Heading2"]))
        for r in REFS:
            story.append(Paragraph("• " + r, s["BodyText"]))
        story.append(PageBreak())
    doc.build(story)

def rasterize(src, dst, dpi=200):
    import pypdfium2 as pdfium
    from PIL import Image
    pdf = pdfium.PdfDocument(src)
    imgs = []
    for i in range(len(pdf)):
        pil = pdf[i].render(scale=dpi / 72).to_pil().convert("L").convert("RGB")
        imgs.append(pil)
    imgs[0].save(dst, save_all=True, append_images=imgs[1:], resolution=dpi)

def build_docx(path):
    import docx
    d = docx.Document()
    d.add_heading(TITLE, 0)
    d.add_paragraph(INTRO)
    for rep in range(6):
        d.add_heading(f"{rep+1}. Werkwijze (deel {rep+1})", level=1)
        for i, (h, body) in enumerate(STEPS, 1):
            d.add_heading(f"{rep+1}.{i} {h}", level=2)
            d.add_paragraph(body)
        d.add_heading("Verantwoordelijkheden", level=2)
        t = d.add_table(rows=0, cols=4); t.style = "Table Grid"
        for row in TABLE:
            cells = t.add_row().cells
            for c, v in zip(cells, row): c.text = v
        d.add_heading("Referenties", level=2)
        for r in REFS: d.add_paragraph(r, style="List Bullet")
    d.save(path)

build_pdf(f"{OUT}/nl_text.pdf")
rasterize(f"{OUT}/nl_text.pdf", f"{OUT}/nl_scan.pdf")
build_docx(f"{OUT}/nl_procedure.docx")
for f in sorted(os.listdir(OUT)):
    print(f, os.path.getsize(os.path.join(OUT, f)))

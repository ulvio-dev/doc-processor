"""Shared fixtures.

The API tests drive the app through httpx's ASGI transport rather than a live
uvicorn, so a test can hold ten SSE streams open at once without ten sockets.
"""
from __future__ import annotations

import io

import pytest

DOCX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


@pytest.fixture(scope="session")
def docx_bytes() -> bytes:
    """A small real .docx, built at test time so no binary lives in the repo."""
    import docx

    d = docx.Document()
    d.add_heading("Quarterly Report", 0)
    d.add_heading("Revenue", level=1)
    d.add_paragraph(
        "Revenue grew 12% to EUR 4.2 million in Q3, driven by the enterprise segment."
    )
    d.add_heading("Costs", level=1)
    d.add_paragraph("Operating costs were flat at EUR 3.1 million.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


@pytest.fixture
def upload(docx_bytes):
    """A multipart `files=` value for the docx, with an optional name override."""

    def _make(name: str = "report.docx", data: bytes | None = None):
        return {"file": (name, data if data is not None else docx_bytes, DOCX_MEDIA_TYPE)}

    return _make


@pytest.fixture
def app_module():
    """The app with its module-level queue state reset.

    The queue, the `_processing` flag and the job registry are module globals
    (one queue per process, by design), so tests must not inherit each other's.
    """
    import src.app as app_mod

    app_mod._queue.clear()
    app_mod._processing = False
    app_mod._jobs.clear()
    yield app_mod
    app_mod._queue.clear()
    app_mod._processing = False
    app_mod._jobs.clear()


@pytest.fixture
def client(app_module):
    import httpx

    transport = httpx.ASGITransport(app=app_module.app)
    return httpx.AsyncClient(transport=transport, base_url="http://test", timeout=120)


@pytest.fixture
def stub_conversion(app_module, monkeypatch):
    """Replace docling with a controllable sleep.

    Queue behaviour (the cap, position countdown, heartbeats) needs jobs slow
    enough to overlap and fast enough for a test suite. Real conversion is
    covered separately in test_process_doc.py.
    """
    import time

    def _install(seconds: float = 0.4, emit_progress: bool = True, fail: bool = False):
        def fake(data, filename, params=None, progress=None):
            if fail:
                raise RuntimeError("boom")
            if emit_progress and progress:
                progress("converting", 5)
            time.sleep(seconds)
            return {
                "markdown": "# stub",
                "chunks": ["stub"],
                "meta": {"filename": filename, "chunk_count": 1, "pages": 1},
            }

        monkeypatch.setattr(app_module, "process_document", fake)

    return _install
